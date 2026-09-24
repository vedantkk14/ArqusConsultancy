"""Bulk import from Excel/CSV: preview vs real run, validation, duplicates, permissions."""

import io

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook

from apps.core.models import AuditLog
from apps.leads.models import Lead

from .conftest import BASE

pytestmark = pytest.mark.django_db
URL = BASE + "/import"
HEAD = ["Name", "Phone", "Email", "Source", "Assigned to", "Requirements", "Proposed value"]


def xlsx(rows, name="leads.xlsx"):
    book = Workbook()
    for row in rows:
        book.active.append(row)
    out = io.BytesIO()
    book.save(out)
    return SimpleUploadedFile(name, out.getvalue())


def post(client, file, **params):
    query = "?" + "&".join(f"{k}={v}" for k, v in params.items()) if params else ""
    return client.post(URL + query, {"file": file}, format="multipart")


def test_import_creates_leads_and_reports(client_for, manager, exec_a):
    rows = [
        HEAD,
        [
            "Rahul Sharma",
            "98765 43210",
            "r@x.com",
            "Referral",
            exec_a.email or exec_a.username,
            "Turf",
            250000,
        ],
        ["Meera", 9876543211, "", "Some Blog", "", "", ""],  # number phone, unknown source
        ["Asha", "12345", "", "", "", "", ""],  # bad phone
        ["", "9876543212", "", "", "", "", ""],  # no name
        ["Dup", "+919876543210", "", "", "", "", ""],  # repeats row 2
    ]
    res = post(client_for(manager), xlsx(rows))
    assert res.status_code == 200
    body = res.json()
    assert body["total"] == 5 and body["created"] == 2 and body["error_count"] == 2
    assert body["duplicate_count"] == 1 and body["duplicates"][0]["row"] == 6
    assert {e["row"] for e in body["errors"]} == {4, 5}
    rahul = Lead.objects.get(phone="+919876543210")
    assert rahul.source == "REFERRAL" and str(rahul.proposed_amount) == "250000.00"
    assert rahul.created_by_id == manager.id
    meera = Lead.objects.get(phone="+919876543211")
    assert meera.source == "OTHER" and meera.source_other == "Some Blog"
    assert AuditLog.objects.filter(model_label="leads.Lead", action="CREATE").count() == 2


def test_assignee_by_name_and_unknown_assignee(client_for, manager, exec_a):
    rows = [
        HEAD,
        ["A", "9876543210", "", "", exec_a.display_name, "", ""],
        ["B", "9876543211", "", "", "Nobody", "", ""],
    ]
    body = post(client_for(manager), xlsx(rows)).json()
    assert body["created"] == 1 and "Nobody" in body["errors"][0]["reason"]
    assert Lead.objects.get(phone="+919876543210").assigned_to_id == exec_a.id


def test_dry_run_creates_nothing(client_for, manager):
    body = post(
        client_for(manager), xlsx([HEAD, ["A", "9876543210"], ["B", "9876543211"]]), dry_run=1
    ).json()
    assert body["dry_run"] is True and body["ready"] == 2 and body["created"] == 0
    assert Lead.objects.count() == 0


def test_existing_leads_are_skipped_not_duplicated(client_for, manager, make_lead):
    make_lead(phone="+919876543210")
    body = post(client_for(manager), xlsx([HEAD, ["A", "9876543210"], ["B", "9876543211"]])).json()
    assert body["created"] == 1 and body["duplicate_count"] == 1
    assert Lead.objects.filter(phone="+919876543210").count() == 1


def test_csv_with_flexible_headers(client_for, manager):
    csv = SimpleUploadedFile(
        "l.csv", "Customer Name,Mobile No\nAnil,98765 43210\n".encode("utf-8-sig")
    )
    assert post(client_for(manager), csv).json()["created"] == 1


@pytest.mark.parametrize(
    "file,message",
    [
        (SimpleUploadedFile("a.pdf", b"x"), "Excel"),
        (SimpleUploadedFile("a.xlsx", b"not a workbook"), "valid .xlsx"),
        (xlsx([["Name", "Email"], ["A", "a@x.com"]]), "Missing column: Phone"),
        (xlsx([]), "no rows"),
    ],
)
def test_unusable_files_are_refused(client_for, manager, file, message):
    res = post(client_for(manager), file)
    assert res.status_code == 400 and message in str(res.json())


def test_row_limit(client_for, manager):
    rows = [HEAD] + [[f"N{i}", "9876543210"] for i in range(1001)]
    res = post(client_for(manager), xlsx(rows))
    assert res.status_code == 400 and "Too many rows" in str(res.json())


def test_only_creators_may_import(client_for, exec_a, pm, manager):
    assert post(client_for(exec_a), xlsx([HEAD, ["A", "9876543210"]])).status_code == 403
    assert post(client_for(pm), xlsx([HEAD, ["A", "9876543210"]])).status_code == 403
    assert client_for(exec_a).get(BASE + "/import-template").status_code == 403
    assert client_for(manager).post(URL, {}, format="multipart").status_code == 400


def test_template_downloads_and_roundtrips(client_for, manager):
    res = client_for(manager).get(BASE + "/import-template")
    assert res.status_code == 200 and "spreadsheetml" in res["Content-Type"]
    sheet = load_workbook(io.BytesIO(res.content)).worksheets[0]
    assert [c.value for c in sheet[1]] == HEAD
    # The example row imports as a real lead (it is meant to be deleted, but must be valid).
    upload = SimpleUploadedFile("t.xlsx", res.content)
    assert post(client_for(manager), upload, dry_run=1).json()["ready"] == 1
