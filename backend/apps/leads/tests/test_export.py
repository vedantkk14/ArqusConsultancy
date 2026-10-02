"""Export as Excel: GET /leads/export?format=xlsx&scope=all|month&month=YYYY-MM."""

import io
from datetime import datetime

import pytest
from openpyxl import load_workbook

from apps.leads import views
from apps.leads.models import Lead, LeadStatus
from apps.leads.selectors import business_tz

from .conftest import BASE


def rows(response) -> list[list]:
    assert response["Content-Type"].startswith(views.XLSX_TYPE)
    book = load_workbook(io.BytesIO(response.content))
    return [list(r) for r in book.active.iter_rows(values_only=True)]


def created_on(lead, year, month, day=15):
    Lead.objects.filter(pk=lead.pk).update(
        created_at=datetime(year, month, day, 12, tzinfo=business_tz())
    )


def test_all_is_every_lead_with_the_current_deal(client_for, manager, make_lead):
    make_lead(name="Asha", status=LeadStatus.WON, proposed_amount=1500)
    make_lead(name="Bala")
    res = client_for(manager).get(f"{BASE}/export?format=xlsx&scope=all")
    assert res.status_code == 200
    assert 'filename="leads-all.xlsx"' in res["Content-Disposition"]
    data = rows(res)
    head = data[0]
    assert head[:5] == ["Name", "Phone", "Email", "Source", "Status"] and "Deals" in head
    by_name = {r[0]: dict(zip(head, r, strict=True)) for r in data[1:]}
    assert set(by_name) == {"Asha", "Bala"}
    assert by_name["Asha"]["Status"] == "WON" and by_name["Asha"]["Proposed value"] == 1500
    assert by_name["Asha"]["Deals"] == 1


def test_xlsx_is_the_default_format(client_for, manager, make_lead):
    make_lead()
    res = client_for(manager).get(f"{BASE}/export")
    assert res.status_code == 200 and len(rows(res)) == 2


def test_month_keeps_only_leads_created_that_month(client_for, admin, make_lead):
    september = make_lead(name="September")
    created_on(september, 2026, 9)
    august = make_lead(name="August")
    created_on(august, 2026, 8, 31)
    res = client_for(admin).get(f"{BASE}/export?format=xlsx&scope=month&month=2026-09")
    assert 'filename="leads-2026-09.xlsx"' in res["Content-Disposition"]
    assert [r[0] for r in rows(res)[1:]] == ["September"]
    res = client_for(admin).get(f"{BASE}/export?format=xlsx&scope=month&month=2026-08")
    assert [r[0] for r in rows(res)[1:]] == ["August"]


def test_month_defaults_to_this_month_and_rejects_nonsense(client_for, admin, make_lead):
    make_lead(name="Now")
    res = client_for(admin).get(f"{BASE}/export?format=xlsx&scope=month")
    assert [r[0] for r in rows(res)[1:]] == ["Now"]
    bad = client_for(admin).get(f"{BASE}/export?format=xlsx&scope=month&month=Sept")
    assert bad.status_code == 400 and "month" in bad.json()["error"]["details"]
    assert client_for(admin).get(f"{BASE}/export?scope=week").status_code == 400


def test_formula_cells_are_neutralised(client_for, manager, make_lead):
    make_lead(name='=HYPERLINK("http://evil")', email="-x@y.co")
    data = rows(client_for(manager).get(f"{BASE}/export?format=xlsx"))
    assert data[1][0] == '\'=HYPERLINK("http://evil")'
    assert data[1][2] == "'-x@y.co"
    assert data[1][1].startswith("'+91")  # phones start with "+"


def test_row_cap(client_for, manager, make_lead, monkeypatch):
    monkeypatch.setattr(views, "EXPORT_MAX_ROWS", 2)
    for _ in range(3):
        make_lead()
    assert len(rows(client_for(manager).get(f"{BASE}/export?format=xlsx"))) == 3  # header + 2


@pytest.mark.parametrize("role", ["exec", "pm"])
def test_only_admin_and_manager_export(client_for, exec_a, pm, role):
    user = {"exec": exec_a, "pm": pm}[role]
    assert client_for(user).get(f"{BASE}/export?format=xlsx").status_code == 403
